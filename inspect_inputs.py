import json
from docx import Document
from openpyxl import load_workbook

doc = Document("assignment_lab3.docx")
payload = {
    "paragraphs": [p.text for p in doc.paragraphs],
    "tables": [
        [[c.text for c in row.cells] for row in table.rows]
        for table in doc.tables
    ],
}
print("DOCX_JSON=" + json.dumps(payload, ensure_ascii=True))

wb = load_workbook("input_lab2.xlsx", data_only=False)
print("SHEETS=" + json.dumps(wb.sheetnames, ensure_ascii=True))
for ws in wb.worksheets:
    meta = {
        "title": ws.title,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "merged": [str(x) for x in ws.merged_cells.ranges],
        "freeze": str(ws.freeze_panes) if ws.freeze_panes else None,
    }
    print("SHEET=" + json.dumps(meta, ensure_ascii=True))
    rows = []
    for row in ws.iter_rows():
        values = [cell.value for cell in row]
        if any(value is not None for value in values):
            rows.append({"r": row[0].row, "v": values})
    print("ROWS=" + json.dumps(rows, ensure_ascii=True, default=str))
